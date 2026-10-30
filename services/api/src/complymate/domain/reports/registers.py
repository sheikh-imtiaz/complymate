"""Register exports: CSV and XLSX from rows of dicts."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any

REPORT_TYPES = {
    "client-register": "Client register",
    "risk-summary": "Risk assessment summary",
    "screening-history": "Screening history",
    "program-versions": "Program version report",
    "training-completion": "Training completion report",
    "audit-trail": "Audit trail report",
}


def _cell(value: Any) -> Any:
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, list | dict):
        return str(value)
    return value


def to_csv(rows: list[dict[str, Any]], columns: list[str] | None = None) -> bytes:
    columns = columns or (list(rows[0].keys()) if rows else [])
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        # Guard against CSV formula injection in spreadsheet apps.
        writer.writerow({k: ("'" + v if isinstance(v, str) and v[:1] in "=+-@" else _cell(v)) for k, v in row.items()})
    return out.getvalue().encode("utf-8-sig")


def to_xlsx(rows: list[dict[str, Any]], title: str, columns: list[str] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    columns = columns or (list(rows[0].keys()) if rows else [])
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F4C81")
    for row in rows:
        ws.append([_cell(row.get(c)) for c in columns])
    for idx, col in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = min(48, max(12, len(col) + 4))
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
