"""Program export: DOCX (python-docx, firm logo + colours) and PDF (WeasyPrint)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

DISCLAIMER = "General information, not legal advice. Have a compliance professional review this program."


@dataclass
class ExportSection:
    title: str
    content_html: str
    citations: list[dict[str, Any]]


@dataclass
class ExportProgram:
    firm_name: str
    version_number: int
    status: str
    approved_at: str | None
    approved_by: str | None
    sections: list[ExportSection]
    logo: bytes | None = None
    primary_color: str = "#0F4C81"


class _HtmlToDocx(HTMLParser):
    """Small HTML → python-docx converter for Tiptap output (p, h2-h4, ul/ol/li, strong, em)."""

    def __init__(self, doc: Any) -> None:
        super().__init__()
        self.doc = doc
        self.para: Any = None
        self.bold = False
        self.italic = False
        self.list_style: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("p", "div"):
            self.para = self.doc.add_paragraph()
        elif tag in ("h2", "h3", "h4"):
            self.para = self.doc.add_heading(level=min(int(tag[1]), 4))
        elif tag == "ul":
            self.list_style.append("List Bullet")
        elif tag == "ol":
            self.list_style.append("List Number")
        elif tag == "li":
            self.para = self.doc.add_paragraph(style=self.list_style[-1] if self.list_style else "List Bullet")
        elif tag in ("strong", "b"):
            self.bold = True
        elif tag in ("em", "i"):
            self.italic = True
        elif tag == "br" and self.para is not None:
            self.para.add_run().add_break()

    def handle_endtag(self, tag: str) -> None:
        if tag in ("ul", "ol") and self.list_style:
            self.list_style.pop()
        elif tag in ("strong", "b"):
            self.bold = False
        elif tag in ("em", "i"):
            self.italic = False
        elif tag in ("p", "li", "h2", "h3", "h4", "div"):
            self.para = None

    def handle_data(self, data: str) -> None:
        if not data.strip():
            return
        if self.para is None:
            self.para = self.doc.add_paragraph()
        run = self.para.add_run(data)
        run.bold = self.bold
        run.italic = self.italic


def to_docx(program: ExportProgram, template_path: str | None = None) -> bytes:
    from docx import Document
    from docx.shared import Inches, RGBColor

    doc = Document(template_path) if template_path and Path(template_path).exists() else Document()
    color = RGBColor.from_string(program.primary_color.lstrip("#").upper()[:6])
    if program.logo:
        doc.add_picture(io.BytesIO(program.logo), width=Inches(1.6))
    title = doc.add_heading(f"{program.firm_name} — AML/CTF Program", level=0)
    for run in title.runs:
        run.font.color.rgb = color
    meta = doc.add_paragraph()
    meta.add_run(f"Version {program.version_number} · Status: {program.status}").bold = True
    if program.approved_at:
        meta.add_run(f"\nApproved {program.approved_at} by {program.approved_by or 'governing body'}")
    meta.add_run(f"\nExported {date.today().isoformat()}")
    for section in program.sections:
        heading = doc.add_heading(section.title, level=1)
        for run in heading.runs:
            run.font.color.rgb = color
        _HtmlToDocx(doc).feed(section.content_html or "<p></p>")
        if section.citations:
            p = doc.add_paragraph()
            p.add_run("Sources: ").italic = True
            p.add_run("; ".join(f"{c.get('title')} ({c.get('url', '')})" for c in section.citations)).italic = True
    doc.add_paragraph().add_run(DISCLAIMER).italic = True
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def to_html(program: ExportProgram) -> str:
    import base64

    from jinja2 import Environment, select_autoescape

    env = Environment(autoescape=select_autoescape(["html"]))
    tpl = env.from_string(
        """<!doctype html><html><head><meta charset="utf-8"><style>
        @page { size: A4; margin: 2cm; @bottom-center { content: "{{ firm }} · AML/CTF Program v{{ v }} · page " counter(page); font-size: 9pt; color: #666; } }
        body { font-family: 'DejaVu Sans', sans-serif; font-size: 10.5pt; line-height: 1.45; color: #1a1a1a; }
        h1 { color: {{ color }}; font-size: 22pt; } h2 { color: {{ color }}; border-bottom: 1px solid #ddd; padding-bottom: 4px; margin-top: 28px; }
        .meta { color: #555; } .sources { font-size: 9pt; color: #555; font-style: italic; } .disclaimer { margin-top: 40px; font-size: 9pt; color: #777; }
        </style></head><body>
        {% if logo %}<img src="data:image/png;base64,{{ logo }}" style="height:56px">{% endif %}
        <h1>{{ firm }} — AML/CTF Program</h1>
        <p class="meta">Version {{ v }} · {{ status }}{% if approved_at %} · Approved {{ approved_at }} by {{ approved_by }}{% endif %}</p>
        {% for s in sections %}<h2>{{ s.title }}</h2>{{ s.content_html | safe }}
        {% if s.citations %}<p class="sources">Sources: {% for c in s.citations %}{{ c.title }}{% if not loop.last %}; {% endif %}{% endfor %}</p>{% endif %}
        {% endfor %}<p class="disclaimer">{{ disclaimer }}</p></body></html>"""
    )
    return tpl.render(
        firm=program.firm_name,
        v=program.version_number,
        status=program.status,
        approved_at=program.approved_at,
        approved_by=program.approved_by,
        sections=[s.__dict__ for s in program.sections],
        color=program.primary_color,
        logo=base64.b64encode(program.logo).decode() if program.logo else None,
        disclaimer=DISCLAIMER,
    )


def html_to_pdf(content: str) -> bytes:
    from weasyprint import HTML

    pdf: bytes = HTML(string=content).write_pdf()
    return pdf


def to_pdf(program: ExportProgram) -> bytes:
    return html_to_pdf(to_html(program))


def safe_filename(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "-", name).strip("-")[:80] or "export"
