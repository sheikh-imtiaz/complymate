"""Downloads sanctions lists (DFAT consolidated list for AU, NZ list), stores a new version when the content
changes and triggers re-screening. Sources are configurable with WATCHLIST_SOURCES (JSON list)."""

from __future__ import annotations

import asyncio
import csv
import io
import json
import os
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import httpx
import structlog

from complymate.container import get_container
from complymate.db.repositories.screening import ScreeningRepo
from complymate.db.session import tenant_session
from complymate.domain.screening.watchlist_parser import checksum, parse
from complymate.observability import metrics
from complymate.ports.queue import QueueName

log = structlog.get_logger(__name__)

DEFAULT_SOURCES = [
    {"source": "DFAT", "url": "https://www.dfat.gov.au/sites/default/files/regulation8_consolidated.xlsx"},
    {
        "source": "NZ",
        "url": "https://www.mfat.govt.nz/assets/Countries-and-Regions/Europe/Ukraine/Russia-Sanctions-Register.xlsx",
    },
]
SAMPLE = Path(__file__).resolve().parents[3] / "seed" / "watchlists_sample.csv"


def xlsx_to_csv(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [[("" if c is None else str(c)).strip() for c in row] for row in ws.iter_rows(values_only=True)]
    header_idx = next((i for i, r in enumerate(rows[:20]) if any("Name" in c for c in r)), 0)
    out = io.StringIO()
    writer = csv.writer(out)
    for r in rows[header_idx:]:
        if any(r):
            writer.writerow(r)
    return out.getvalue()


async def _fetch(url: str) -> str:
    async with httpx.AsyncClient(timeout=60, follow_redirects=True) as http:
        resp = await http.get(url, headers={"User-Agent": "ComplyMate watchlist refresh"})
        resp.raise_for_status()
        if url.endswith(".xlsx") or resp.headers.get("content-type", "").startswith("application/vnd.openxml"):
            return xlsx_to_csv(resp.content)
        return resp.text


async def load_content(source: dict[str, str]) -> str:
    if source.get("url", "").startswith("file://") or not source.get("url"):
        path = Path(source.get("url", f"file://{SAMPLE}").removeprefix("file://"))
        return await asyncio.to_thread(path.read_text, encoding="utf-8")
    return await _fetch(source["url"])


async def run(payload: dict[str, Any]) -> dict[str, Any]:
    container = get_container()
    sources = json.loads(os.environ.get("WATCHLIST_SOURCES", "null")) or (
        [{"source": "SAMPLE", "url": f"file://{SAMPLE}"}] if container.settings.is_local else DEFAULT_SOURCES
    )
    changed: list[str] = []
    for src in sources:
        try:
            content = await load_content(src)
            entries = parse(content, src["source"])
            if not entries:
                raise ValueError("no entries parsed")
            async with tenant_session(None, platform=True) as session:
                version = await ScreeningRepo(session).replace_watchlist(
                    src["source"], f"{src['source']} {datetime.now(UTC):%Y-%m-%d}", checksum(content), entries
                )
            if version is not None:
                changed.append(src["source"])
                log.info("watchlist.updated", source=src["source"], entries=len(entries))
        except Exception:
            metrics.screening_job_failed(f"refresh_watchlists:{src['source']}")
            log.exception("watchlist.refresh_failed", source=src["source"])
    if changed and not payload.get("skip_rescreen"):
        await container.queue.send(
            QueueName.SCREENING, {"task": "nightly_rescreen", "payload": {"reason": "lists_updated"}}
        )
    return {"changed": changed}
